import mysql.connector
from datetime import datetime
import openpyxl
import re
 
file_path = r"C:\Users\lynette.mutuku\Documents\ett3\OpCO - Jan.xlsx"
sheet_name = "Sheet2"

matrix = [ 
['Bayobab KE','@Bayobab.','Bayobab KE'],
['MTN Uganda','.ug@','Uganda' , ],
['MTN Nigeria','.NG@','Nigeria'],
['MTN Ghana','.GH@','Ghana'],
['MTN Rwanda','.rw@','Rwanda'],
['MTN Zambia','.zm@','Zambia'],
['MTN Cameroon','.cm@','Cameroon'],
['MTN South Sudan','nmc.SS','South Sudan'],
['MTN South Sudan','NOC.SS','South Sudan'],
['MTN Guinea Conakry','.GN@','GUINEA-REPUBLIC'],
['MTN Guinea Bissau','.gw@','Guinea-Bissau'],
['MTN Eswatini','.SZ@','Eswatini'],
['MTN Botswana','.co.bw','Botswana'],
['MTN Liberia','mtnnocfolib@huawei.com','Liberia'],
['MTN Liberia','.LR@','Liberia'],
['MTN CIV','.mtnci@','Cote '],
['MTN South Africa','.za@','mtn sa'],
['MTN South Africa','.mtnsa@','mtn sa'],
['MTN South Africa','wholesaleza@','mtn sa'],
['MTN Afghanistan','.AF@','Afghanistan'],
['MTN Congo','nocmtncg','Congo'],
['MTN Benin','.bj@','Benin'],
['MTN Sudan','nmc@mtn.sd','Sudan'],
]

# matrix2 = [
#     ['Bayobab KE','@Bayobab.','Bayobab KE'],
#     ['MTN Zambia','.zm@','Zambia'],
# ]
 
 
 
def getmttr(sendTime , responsereceiveTime):
 
    # Parse the timestamps into datetime objects
    dt1 = datetime.strptime(sendTime, '%Y-%m-%d %H:%M:%S')
    dt2 = datetime.strptime(responsereceiveTime, '%Y-%m-%d %H:%M:%S')
 
    # Calculate the time difference in minutes
    diff_minutes = round((dt2 - dt1).total_seconds() / 60)
    return diff_minutes
    #print(f"Difference in minutes: {diff_minutes}")
 
 
def append_data_to_excel(file_path, sheet_name, data):
    # Load the workbook
    workbook = openpyxl.load_workbook(file_path)
 
    # Select the sheet
    sheet = workbook[sheet_name]
 
    # Find the last used row
    last_row = sheet.max_row
 
    # Append the data to the next row
    next_row = last_row + 1
    for i, value in enumerate(data, start=1):
        sheet.cell(row=next_row, column=i).value = value
 
    # Save the workbook
    workbook.save(file_path)
 
# Establish a connection to the MySQL database
cnx = mysql.connector.connect(
host='127.0.0.1',
# user='vicky',
# password='ILoveBlack@2022',
user='lyn',
password='0gDUnc55',
database='CSCMailsBackup',
charset='utf8mb4',
collation='utf8mb4_general_ci',
auth_plugin='mysql_native_password'
)
 
# Create a cursor object to execute SQL queries
cursor = cnx.cursor()
# Define your SQL query
 
 
for i in range(0 , len(matrix)):
 
    OPCO = matrix[i][0]
    partner = matrix[i][1]
    partner2 = matrix[i][2]
   
 
    print("\n")
    print("---------------------------------------------------------------------------------------")
    print("To Domain: " , partner)
 
    if i==0 :
    #     sql_query = '''SELECT * FROM Mails WHERE YEAR(MailDate) = 2025 AND MONTH(MailDate) = 12 AND MailCategory = 'Sent'
    # AND IncidentNo NOT LIKE 'RITM%' AND
    # (EmailAddress like '%.ke@%'  or EmailAddress like '%MTN Support [ Bayobab Kenya ]%') and
    # IncidentNo != ''  AND
    # MailType = 'Notification' order by MailDate desc;'''
 
        sql_query = '''SELECT * FROM Mails WHERE YEAR(MailDate) = 2025 AND MONTH(MailDate) = 12 AND MailCategory = 'Sent'
    AND MailSubject NOT LIKE '%RITM%' 
    AND (EmailAddress like '%.ke@%'  or EmailAddress like '%MTN Support [ Bayobab Kenya ]%') 
    AND MailSubject NOT LIKE 'Re:%'
    AND MailSubject NOT LIKE 'Escalation%'
    AND MailType = 'Notification' 
    order by MailDate desc;'''
 
    else:
        # Previous code that's not working
       
        # sql_query = '''SELECT * FROM Mails WHERE YEAR(MailDate) = 2025 AND MONTH(MailDate) = 12 AND MailCategory = 'Sent'
        # AND IncidentNo NOT LIKE 'RITM%' AND
        # (EmailAddress like '%'''+partner+'''%'  or EmailAddress like '%'''+partner2+'''%') and
        # IncidentNo != ''  AND
        # MailType = 'Notification' order by MailDate desc;'''
        
        sql_query = '''SELECT * FROM Mails WHERE YEAR(MailDate) = 2025 AND MONTH(MailDate) = 12 AND MailCategory = 'Sent'
        AND MailSubject NOT LIKE '%RITM%'
        AND (EmailAddress like '%'''+partner+'''%'  or EmailAddress like '%'''+partner2+'''%') 
        AND MailSubject NOT LIKE 'Re:%' 
        AND MailSubject NOT LIKE 'Escalation%'
        AND MailType = 'Notification' 
        order by MailDate desc;'''
 
 
    cursor.execute(sql_query)
 
    # Fetch All row returned by the query
    row = cursor.fetchall()
    print("Total Items Raised : " , len(row))
 
    i = 1
    for item in row:  

        # print("RAW ITEM!!!!!!!!!!: ", item)    
 
        subject = str(item[1])
        addresses = str(item[5])
        names = str(item[6])
        # incidentNo = str(item[3])  

        match = re.search(r'\bINC\w+', subject)

        incidentNo = match.group() if match else None

        if not incidentNo:
            print("Skipping row — no INC found:", subject)
            continue

        print("INCIDENT NO", incidentNo)  
 
        #print(subject , "\n")
        #print(addresses)
        #print(names)
        #print("---------------------------------------------------------------------------------------")

 
        if OPCO == 'Bayobab KE':
            sql_query = (
                "SELECT * FROM Mails  "
                "WHERE MailCategory = 'Inbox'  "
                "AND MailSubject LIKE '%" + incidentNo + "%'  "
                "AND (EmailName LIKE '%Mercy Mulondu%' "
                "OR EmailName LIKE '%Hilda Nasambu%'"
                "OR EmailName LIKE '%Alfred Lunalo%'"
                "OR EmailName LIKE '%Rose Naomi%'"
                "OR EmailName LIKE '%Regina Mutuku%'"
                "OR EmailName LIKE '%Stephanie Iguanya%'"
                "OR EmailName LIKE '%Chris Mutwiri%')"
                "ORDER BY MailDate ASC LIMIT 1")
        else:           
            sql_query = '''SELECT * FROM Mails WHERE MailCategory = 'Inbox' 
            AND MailSubject LIKE '%" + incidentNo + "%' 
            AND (EmailAddress LIKE '%" + partner + "%' or EmailName LIKE '%" + partner2 + "%' ) 
            ORDER BY MailDate ASC LIMIT 1'''

        # print(sql_query)
        cursor.execute(sql_query)
        response = cursor.fetchall()  

        if not len(response) == 0:
 
            sendTime = str(item[2])
            responsereceiveTime = str(response[0][2])
 
            mttr = getmttr(sendTime , responsereceiveTime)
 
            if mttr > 0:
 
                print(i ,".", OPCO ,  incidentNo , ":" , item[2] ,  "<-->" , response[0][2] , "MTTR = " , mttr  )
 
                data = [i, OPCO ,incidentNo, item[1] , sendTime  , response[0][1] , responsereceiveTime ,  mttr , "Acknowledged"]
                append_data_to_excel(file_path, sheet_name, data)
 
        else:
 
            print(i ,".", OPCO ,incidentNo , ":" , item[2] ,  "<-->" , "No response"  )
 
            data = [i, OPCO , incidentNo , item[1] ,item[2] , "No Response" , "No Response" , "No Response", "Not Acknowledged"]
 
            append_data_to_excel(file_path, sheet_name, data)
       
       
        i = i + 1
       
 
    #print(response)
   