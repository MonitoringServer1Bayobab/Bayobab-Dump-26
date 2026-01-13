import mysql.connector
from datetime import datetime
import openpyxl
import re
 
file_path = r"C:\Users\lynette.mutuku\Documents\ett3\3rdParty\3rdPartyAnalysisDec.xlsx"
sheet_name = "Sheet1"
 
partner = ""
partner2 = ""
 
partnersDB = {'ucom.mobi' : 'Ucom',
'orange-sonatel.com' : 'ACE',
'guilab.com.gn' : 'ACE',
'africadatacentres.com' : 'ADC',
'africell.cd' : 'Africell',
'afr-' : 'Afr-Ix',
'agiletelecom.com' : 'Agile Telecom',
'ng.airtel.com' : 'Airtel',
'airtel.com' : 'Airtel',
'mw.airtel.com' : 'Airtel Malawi',
'amazon.com' : 'Amazon',
'ams-ix.net' : 'Ams Ix',
'anam.com' : 'Anam Support',
'intl.att.com' : 'AT&T',
'att.com' : 'AT&T',
'snoc-portal.atlassian.net' : 'Atlassian',
'axione.com' : 'Axione',
# 'ayoba.me' : 'Ayoba',
'bics.com' : 'BICS',
'blackberry.com' : 'Blackberry',
'bofinet.co.bw' : 'Bofinet',
'intellico.ch' : 'Calltrade',
'panaceamobile.com' : 'Cellfind',
'cellusys.com' : 'Cellusys (Firewall Team)',
'cequens.com' : 'Cequens',
'chinatelecomglobal.com' : 'China Telecom',
'cm.com' : 'Cm Telecom',
'cmcnetworks.net' : 'CMC',
'cmi.chinamobile.com' : 'CMI',
'cogentco.com' : 'Cogent',
'control.center@colt.net' : 'Colt',
'colt.ne' : 'Colt',
'telekom.de' : 'Deutsche Telekom',
'intnet.dj' : 'Djibouti Telecom',
'djibtel.dj' : 'Djibouti Telecom',
'dolphintelecom.net' : 'Dolphin Ghana',
'du.ae' : 'Du Telecom',
'openserve.co.za' : 'Eassy',
'europeindiagateway.com' : 'EIG',
'epsilontel.com' : 'Epsilon',
'ethiotelecom.et' : 'Ethiotelecom',
'etisalat.ae' : 'Etisalat',
'eand.com' : 'Etisalat',
'exchangetelecom.com.ng' : 'Exchange',
'expresso-global.com' : 'Express Telecom Group',
'expressotelecom.com' : 'Expresso Telecom Group Escalation Matrix',
'fb.com' : 'Facebook',
'netsf.fr' : 'Free senegal',
'freebusiness.sn' : 'Free Senegal',
'javna.com' : 'Friendly',
'gabsgroup.com' : 'GABS',
'gmail.com' : 'Globacom',
'globalswitch.com' : 'Global Switch',
'globaltechnics' : 'Global Technics',
'glo1' : 'Gloworld',
'gms-worldwide.com' : 'Gms A2P',
'globeteleservices.com' : 'GTS',
'gtt.net' : 'GTT',
'gva.africa' : 'GVA',
'canalbox.net' : 'GVA',
'ibasis.net' : 'Ibasis',
'imimobile.com' : 'IMI',
'infobip.com' : 'Infobip',
'ipinfotech.com' : 'Ip Infotech',
'isoceltelecom.com' : 'Isocel Sa',
'jazz.com.pk' : 'JAZZ',
'jenny' : 'Jenny',
'jtl.co.ke' : 'JTL',
'kaleyra.com' : 'Kaleyra',
'ventatele.com' : 'Lanck Telecom',
'horisen-messaging.com' : 'Link Mobility',
'liquid.tech' : 'Liquid Kenya',
'liquidtelecom.com' : 'Liquid Zambia',
'lleida.net' : 'Lleida',
'lmtgroup.com' : 'LMT',
'plintron.com' : 'LYCA',
'mainone.net' : 'Mainone',
'iam.ma' : 'Maroc',
'mascom.bw' : 'Mascom',
'eljawal.mr' : 'Mauritel',
'altice.pt' : 'MEO',
'telecom.pt' : 'Meo-Ptc',
'cprm.net' : 'Meo-Ptc',
'messagebird.com' : 'Messagebird',
'tanla.com' : 'Mgage',
'karix.com' : 'Mgage',
'mgi-management.com' : 'MGI',
'mitto.ch' : 'Mitto',
'mobileum.com' : 'Mobelium',
'mobily.com.sa' : 'Mobily',
'montymobile.com' : 'Monty Mobile',
'moyanetworks.com' : 'Moya',
'nexmo.com' : 'Nexmo (Vonage)',
'nokia.com' : 'Npo Lead',
'omantel.om' : 'Omantel',
'openmarket.com' : 'Open Market',
'orange.com' : 'Orange',
'csciw@orange.com' : 'Orange International',
'oci@orange.com' : 'Orange International Circuits',
'paratus.ao' : 'Paratus',
'paratus.africa' : 'Paratus',
'ita.ao' : 'Paratus',
'pccwglobal.com' : 'PCCW',
'ptcl.net.pk' : 'PTCL',
'realnet.co.sz' : 'Real Image',
'routemobile.com' : 'Routemobile',
'safaricom.co.ke' : 'Safaricom Escalation List',
'salcab.sl' : 'Salcab',
'sama-telecom.com' : 'Sama Telecom',
'sinch.com' : 'SAP',
'sbin.bj' : 'SBIN',
'seacom.mu' : 'Seacom',
'seacom.com' : 'Seacom',
'seacom' : 'Seacom',
'pamoja-africa.com' : 'Seacom',
'siliconeconnect.com' : 'Silicon Connect',
'smshighway.com' : 'Sms Highway',
'tisparkle.com' : 'Sparkle',
'stc.com.sa' : 'STC',
'sudatel' : 'Sudatel',
'syniverse.com' : 'Syniverse',
'tatacommunications.com' : 'Tata',
'telecelglobal.com' : 'Telecel Global Information',
'telecity.com' : 'Telecity',
'digitalrealty.com' : 'Telecity',
'@tm.com.my' : 'Telecom Malaysia',
'telecom.na' : 'Telecom Namibia',
'telefonica.com' : 'Telefonica',
'uk.telehouse.net' : 'Telehouse North',
# 'telkom.co.ke' : 'Telkom Kenya',
'nmc@Telkom.co.ke' : 'TKL',
'telkom.co.za' : 'Telkom SA',
'tigo.co.tz' : 'TIGO',
'yas.co.tz' : 'TIGO',
'tmcel.mz' : 'Tmcel',
'togotelecom.tg' : 'Togocom',
'ttcl.co.tz' : 'TTCL',
'twilio.com' : 'Twilio',
'tyntec.com' : 'Tyntec',
'ucom.mobi' : 'Ucom',
'viber.com' : 'Viber',
'viettel.com.vn' : 'Viettel',
'vodacom.co.za' : 'Vodacom',
'vodafone.com' : 'Vodafone',
'vodatel.com' : 'Vodatel',
'voxcarrier.com' : 'Vox',
'liquidtelecom.co.za' : 'WACS',
'infraco.co.za' : 'WACS',
'wiocc.net' : 'Wiocc',
'wis.one' : 'Wis Escalation List',
'wtl.co.ke' : 'WTL',
'sa.zain.com' : 'Zain Ksa',
'zamtel.co.zm' : 'Zamtel (Zambia Telecommunications Company)',
'zantel.co.tz' : 'Zantel',
'zayo.com' : 'Zayo',
'zesco.co.zm' : 'Zesco',
'equinix.com':'Equinix',
'pccwglobal.com':'PCCW',
'infraco.co.za':'BBI',
'mtc.com.na':'MTC'
}    
 
currentList = ['Ucom',
'ACE',
'ACE',
'ADC',
'Africell',
'Afr-Ix',
'Agile Telecom',
'Airtel',
'Airtel',
'Airtel Malawi',
'Amazon',
'Ams Ix',
'Anam Support',
'AT&T',
'AT&T',
'Atlassian',
'Axione',
# 'Ayoba',
'BICS',
'Blackberry',
'Bofinet',
'Calltrade',
'Cellfind',
'Cellusys (Firewall Team)',
'Cequens',
'China Telecom',
'Cm Telecom',
'CMC',
'CMI',
'Cogent',
'Colt',
'Colt',
'Deutsche Telekom',
'Djibouti Telecom',
'Djibouti Telecom',
'Dolphin Ghana',
'Du Telecom',
'Eassy',
'EIG',
'Epsilon',
'Ethiotelecom',
'Etisalat',
'Etisalat',
'Exchange',
'Express Telecom Group',
'Expresso Telecom Group Escalation Matrix',
'Facebook',
'Free senegal',
'Free Senegal',
'Friendly',
'GABS',
'Globacom',
'Global Switch',
'Global Technics',
'Gloworld',
'Gms A2P',
'GTS',
'GTT',
'GVA',
'GVA',
'Ibasis',
'IMI',
'Infobip',
'Ip Infotech',
'Isocel Sa',
'JAZZ',
'Jenny',
'JTL',
'Kaleyra',
'Lanck Telecom',
'Link Mobility',
'Liquid Kenya',
'Liquid Zambia',
'Lleida',
'LMT',
'LYCA',
'Mainone',
'Maroc',
'Mascom',
'Mauritel',
'MEO',
'Meo-Ptc',
'Meo-Ptc',
'Messagebird',
'Mgage',
'Mgage',
'MGI',
'Mitto',
'Mobelium',
'Mobily',
'Monty Mobile',
'Moya',
'Nexmo (Vonage)',
'Npo Lead',
'Omantel',
'Open Market',
'Orange',
'Orange International',
'Orange International Circuits',
'Paratus',
'Paratus',
'Paratus',
'PCCW',
'PTCL',
'Real Image',
'Routemobile',
'Safaricom Escalation List',
'Salcab',
'Sama Telecom',
'SAP',
'SBIN',
'Seacom',
'Seacom',
'Seacom',
'Seacom',
'Silicon Connect',
'Sms Highway',
'Sparkle',
'STC',
'Sudatel',
'Syniverse',
'Tata',
'Telecel Global Information',
'Telecity',
'Telecity',
'Telecom Malaysia',
'Telecom Namibia',
'Telefonica',
'Telehouse North',
# 'Telkom Kenya',
'TKL',
'Telkom SA',
'TIGO',
'Tmcel',
'Togocom',
'TTCL',
'Twilio',
'Tyntec',
'Ucom',
'Viber',
'Viettel',
'Vodacom',
'Vodafone',
'Vodatel',
'Vox',
'WACS',
'WACS',
'Wiocc',
'Wis Escalation List',
'WTL',
'Zain Ksa',
'Zamtel (Zambia Telecommunications Company)',
'Zantel',
'Zayo',
'Zesco',
'Equinix',
'PCCW',
'BBI',
'MTC'
]
 
partnersDB2 = {'nmc@Telkom.co.ke' : 'TKL'}
 
currentList2 = ['TKL',]
# partnersDB2 = {'@tm.com.my' : 'Telecom Malaysia',
#                'intnet.dj' : 'Djibouti Telecom',
# 'djibtel.dj' : 'Djibouti Telecom',
#                }
 
# currentList2 = [
#     'Telkom Malaysia',
#     'Djibouti Telecom',
#     'Djibouti Telecom',
#                 ]
 
 
def getmttr(sendTime , responsereceiveTime):
 
    # Parse the timestamps into datetime objects
    dt1 = datetime.strptime(sendTime, '%Y-%m-%d %H:%M:%S')
    dt2 = datetime.strptime(responsereceiveTime, '%Y-%m-%d %H:%M:%S')
 
    # Calculate the time difference in minutes
    diff_minutes = round((dt2 - dt1).total_seconds() / 60)
    return diff_minutes
    #print(f"Difference in minutes: {diff_minutes}")
 


def clean_illegal_chars(value):
    if isinstance(value, str):
        # Remove control characters (0x00-0x1F except \t, \n, \r)
        return re.sub(r'[\x00-\x08\x0B\x0C\x0E-\x1F]', '', value)
    return value

 
def append_data_to_excel(file_path, sheet_name, data):
    # Load the workbook
    workbook = openpyxl.load_workbook(file_path)
 
    # Select the sheet
    sheet = workbook[sheet_name]
 
    # Find the last used row
    last_row = sheet.max_row
 
    # Append the data to the next row
    next_row = last_row + 1
    for i, value in enumerate(data, start=1):
        sheet.cell(row=next_row, column=i).value = value
 
    # Save the workbook
    workbook.save(file_path)
 
# # Establish a connection to the MySQL database
# cnx = mysql.connector.connect(
# host='127.0.0.1',
# user='Dennis',
# password='Allan2020*',
# database='CSCMailsBackup'
# )

cnx = mysql.connector.connect(
host='127.0.0.1',
# user='vicky',
# password='ILoveBlack@2022',
user='lyn',
password='0gDUnc55',
database='CSCMailsBackup',
charset='utf8mb4',
collation='utf8mb4_general_ci',
auth_plugin='mysql_native_password'
)
 
# Create a cursor object to execute SQL queries
cursor = cnx.cursor()
# Define your SQL query
 
 
for k in range(11 , len(currentList)):
    #get user input
    #partner = input("Please type a partner: ")
    #partner2 = input("Please type an alias: ")
 
    partner = str(currentList[k])
    partner2 = partner
 
    print("\n")
    print("---------------------------------------------------------------------------------------")
    print("To Domain: " , partner)
    # print("Partner2: " , partner2)

    # break
 
    sql_query = (
        "SELECT * FROM Mails WHERE YEAR(MailDate) = 2025 AND Month(MailDate) = 12 AND MailCategory = 'Sent' "
        "AND (EmailAddress LIKE '%" + partner + "%' or EmailName LIKE '%" + partner2 +    "%' "
        "or EmailName LIKE '%" + partner + "%' or EmailAddress LIKE '%" + partner2 + "%' ) "
        "AND IncidentNo NOT LIKE 'RITM%' "
        "AND IncidentNo != ''  AND MailType = 'Notification' order by MailDate desc"
    )
 
 
    # Execute the SQL query
    cursor.execute(sql_query)
 
    # Fetch All row returned by the query
    row = cursor.fetchall()
    print("Total Items Raised : " , len(row))
 
    i = 1
    for item in row:      
 
        partner = str(currentList[k])
        partner2 = partner
 
        incidentNo = str(item[3])      
       
        sql_query = "SELECT * FROM Mails WHERE MailCategory = 'Inbox' AND MailSubject LIKE '%" + incidentNo + "%' AND (EmailAddress LIKE '%" + partner + "%' or EmailName LIKE '%" + partner2 + "%' ) ORDER BY MailDate ASC LIMIT 1"
        cursor.execute(sql_query)
        response = cursor.fetchall()      
     
        partner = partnersDB.get(partner)
        # Checking whether the list size is equal to 0
        if not len(response) == 0:
       
            sendTime = str(item[2])
            responsereceiveTime = str(response[0][2])
 
            mttr = getmttr(sendTime , responsereceiveTime)
 
            if mttr > 0:
 
                print(i ,".", partner ,  item[3] , ":" , item[2] ,  "<-->" , response[0][2] , "MTTR = " , mttr  )
 
                # data = [i, partner ,item[3], item[1] , sendTime  , response[0][1] , responsereceiveTime ,  mttr ,"Acknowledged", partner2]
                # append_data_to_excel(file_path, sheet_name, data)
                data = [i, partner, item[3], item[1], sendTime, response[0][1], responsereceiveTime, mttr, "Acknowledged", partner2]
                cleaned_data = [clean_illegal_chars(v) for v in data]
                append_data_to_excel(file_path, sheet_name, cleaned_data)
 
        else:
       
            print(i ,".", partner ,item[3] , ":" , item[2] ,  "<-->" , "No response"  )
 
            # data = [i, partner , item[3] , item[1] ,item[2] , "No Response" , "No Response" , "No Response", "Not Acknowledged" , partner2]
 
            # append_data_to_excel(file_path, sheet_name, data)
            data = [i, partner, item[3], item[1], item[2], "No Response", "No Response", "No Response", "Not Acknowledged", partner2]
            cleaned_data = [clean_illegal_chars(v) for v in data]
            append_data_to_excel(file_path, sheet_name, cleaned_data)
            pass
 
        i = i + 1